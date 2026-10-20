





from flask import Flask, render_template, request, redirect, url_for, session, flash, Response, jsonify
import sqlite3, time, os, uuid
from dotenv import load_dotenv
load_dotenv()
from werkzeug.utils import secure_filename

try:
    from google import genai
    _use_new_genai = True
except ImportError:
    import google.generativeai as genai
    _use_new_genai = False

app = Flask(__name__, template_folder="frontend/templates", static_folder="frontend/static")
app.secret_key = "super_secret_key"

# ========== FILE UPLOAD CONFIG ==========
UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), "instance", "uploads")
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt", "md", "csv", "py", "js", "html", "json", "png", "jpg", "jpeg", "gif", "bmp", "webp"}
IMAGE_MIME = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
              "gif": "image/gif", "bmp": "image/bmp", "webp": "image/webp"}
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # 20 MB limit
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_file(filepath, filename):
    """Extract readable text from uploaded file."""
    ext = filename.rsplit('.', 1)[1].lower()
    text = ""
    try:
        if ext == "pdf":
            try:
                import pdfplumber
                with pdfplumber.open(filepath) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
            except ImportError:
                try:
                    import PyPDF2
                    with open(filepath, "rb") as f:
                        reader = PyPDF2.PdfReader(f)
                        for page in reader.pages:
                            text += (page.extract_text() or "") + "\n"
                except ImportError:
                    text = "[PDF text extraction requires pdfplumber or PyPDF2. Please install one: pip install pdfplumber]"
        elif ext == "docx":
            try:
                import docx
                d = docx.Document(filepath)
                text = "\n".join(p.text for p in d.paragraphs)
                for t in d.tables:
                    for row in t.rows:
                        text += "\n" + " | ".join(c.text for c in row.cells)
            except ImportError:
                text = "[DOCX extraction requires python-docx. Install: pip install python-docx]"
        elif ext in {"txt", "md", "csv", "py", "js", "html", "json"}:
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        elif ext in {"png", "jpg", "jpeg", "gif", "bmp", "webp"}:
            text = ""  # images are sent to the model directly (see /ask-with-file)
    except Exception as e:
        text = f"[Could not extract text: {e}]"
    return text.strip()

# ========== SET GEMINI API KEY ==========
API_KEY = os.environ.get("GEMINI_API_KEY", "")  # 🔹 Set GEMINI_API_KEY in your .env file
MODEL_ID = "gemini-3.8-flash"

if _use_new_genai:
    client = genai.Client(api_key=API_KEY)
    model = None
else:
    genai.configure(api_key=API_KEY)
    model = genai.GenerativeModel(MODEL_ID)
    client = None

# ================== DATABASE INIT ==================
def init_db():
    con = sqlite3.connect("instance/database.db")
    cur = con.cursor()
    cur.execute("CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT, password TEXT)")
    con.commit()
    con.close()

init_db()

# ================== ROUTES ==================

@app.route('/')
def index():
    return render_template("index.html")

@app.route('/signup', methods=['GET','POST'])
def signup():
    if request.method == "POST":
        username = request.form['username']
        password = request.form['password']
        con = sqlite3.connect("instance/database.db")
        cur = con.cursor()
        cur.execute("INSERT INTO users (username,password) VALUES (?,?)",(username,password))
        con.commit()
        con.close()
        flash("Account Created! Please login.")
        return redirect(url_for("login"))
    return render_template("signup.html")

@app.route('/login', methods=['GET','POST'])
def login():
    if request.method == "POST":
        username = request.form['username']
        password = request.form['password']
        con = sqlite3.connect("instance/database.db")
        cur = con.cursor()
        cur.execute("SELECT * FROM users WHERE username=? AND password=?",(username,password))
        user = cur.fetchone()
        con.close()
        if user:
            session['username'] = username
            return redirect(url_for("dashboard"))
        else:
            flash("Invalid Credentials")
    return render_template("login.html")

@app.route('/dashboard')
def dashboard():
    if "username" not in session:
        return redirect(url_for("login"))
    return render_template("dashboard.html", user=session['username'])

@app.route('/chat')
def chat():
    if "username" not in session:
        return redirect(url_for("login"))
    return render_template("chat.html", user=session['username'])

@app.route('/ask', methods=['POST'])
def ask_ai():
    question = request.form['question']

    def generate():
        prompt_text = f"""
            Explain the following clearly in **Markdown format**. 
            - Use headings (##), bold, and underline where needed.  
            - Use tables if needed.  
            - Do NOT wrap in HTML tags, only markdown.  

            Question: {question}
            """
        try:
            if _use_new_genai:
                response = client.models.generate_content_stream(
                    model=MODEL_ID,
                    contents=prompt_text
                )
            else:
                response = model.generate_content(prompt_text, stream=True)

            for chunk in response:
                text = getattr(chunk, 'text', None)
                if text:
                    yield text
                    time.sleep(0.01)
        except Exception as e:
            yield f"\n\n**Error:** {e}"

    return Response(generate(), mimetype='text/plain')


@app.route('/ask-with-file', methods=['POST'])
def ask_ai_with_file():
    """Handle chat with optional file attachments."""
    if "username" not in session:
        return Response("\n\n**Session expired.** Please log in again.", status=401, mimetype='text/plain')
    question = request.form.get('question', '').strip()
    uploaded_files = request.files.getlist('files')

    file_contents = []
    image_parts = []
    for file in uploaded_files:
        if file and file.filename and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            ext = filename.rsplit('.', 1)[1].lower()
            if ext in IMAGE_MIME:
                image_parts.append((filename, IMAGE_MIME[ext], file.read()))
                continue
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], f"{uuid.uuid4().hex}_{filename}")
            file.save(filepath)
            extracted = extract_text_from_file(filepath, filename)
            if extracted:
                file_contents.append(f"=== File: {filename} ===\n{extracted}")
            try:
                os.remove(filepath)
            except Exception:
                pass

    def generate():
        if file_contents or image_parts:
            files_section = "\n\n".join(file_contents) if file_contents else "(No text files. See the attached image(s): " + ", ".join(n for n, _, _ in image_parts) + ")"
            prompt_text = f"""
You are KNOWZA, an intelligent exam preparation assistant. The user has uploaded the following file(s) for context:

{files_section}

---

Based on the above content, answer the following question clearly in **Markdown format**:
- Use headings (##), bold, bullet points, and tables where appropriate.
- Do NOT wrap in HTML tags, only markdown.

Question: {question if question else 'Summarize the uploaded file(s)/image(s) and highlight key concepts useful for exam preparation.'}
"""
        else:
            prompt_text = f"""
Explain the following clearly in **Markdown format**.
- Use headings (##), bold, and bullet points where needed.
- Use tables if needed.
- Do NOT wrap in HTML tags, only markdown.

Question: {question}
"""
        try:
            contents = prompt_text
            if image_parts:
                if _use_new_genai:
                    from google.genai import types
                    contents = [prompt_text] + [types.Part.from_bytes(data=d, mime_type=m) for _, m, d in image_parts]
                else:
                    contents = [prompt_text] + [{"mime_type": m, "data": d} for _, m, d in image_parts]
            if _use_new_genai:
                response = client.models.generate_content_stream(
                    model=MODEL_ID,
                    contents=contents
                )
            else:
                response = model.generate_content(contents, stream=True)

            for chunk in response:
                text = getattr(chunk, 'text', None)
                if text:
                    yield text
                    time.sleep(0.01)
        except Exception as e:
            yield f"\n\n**Error:** {e}"

    return Response(generate(), mimetype='text/plain')

@app.errorhandler(413)
def too_large(e):
    return Response("\n\n**File too large.** Total upload limit is 20 MB.", status=413, mimetype='text/plain')

# ✅ Auto Quiz Generator API for frontend JS
def _parse_quiz_json(raw):
    """Parse the model's JSON quiz; returns a clean list or []."""
    import json, re
    txt = raw.strip()
    txt = re.sub(r'^```(?:json)?\s*|\s*```$', '', txt, flags=re.IGNORECASE)
    m = re.search(r'\[.*\]', txt, re.DOTALL)
    if not m:
        return []
    try:
        items = json.loads(m.group(0))
    except Exception:
        return []
    out = []
    for it in items:
        try:
            opts = [re.sub(r'^[A-D][\.\)]\s*', '', str(o).strip()) for o in it["options"]]
            ans = str(it["answer"]).strip().upper()[:1]
            if len(opts) == 4 and ans in "ABCD" and ans:
                out.append({"question": str(it["question"]).strip(), "options": opts,
                            "answer": ans, "explanation": str(it.get("explanation", "")).strip()})
        except Exception:
            continue
    return out

@app.route('/api/generate-quiz', methods=['POST'])
def api_generate_quiz():
    if "username" not in session:
        return jsonify({"error": "Please log in again."}), 401
    data = request.get_json() or {}
    topic = str(data.get("topic", "General Knowledge")).strip() or "General Knowledge"
    try:
        num = max(1, min(30, int(data.get("numQuestions", 5))))
    except (TypeError, ValueError):
        num = 5
    difficulty = str(data.get("difficulty", "medium")).strip().lower()
    if difficulty not in ("easy", "medium", "hard"):
        difficulty = "medium"

    guide = {
        "easy": "basic recall and definitions, suitable for beginners",
        "medium": "moderate understanding and application of concepts",
        "hard": "advanced, tricky questions needing deep understanding and multi-step reasoning",
    }[difficulty]

    prompt = f"""Generate exactly {num} multiple-choice questions on the topic: {topic}.
Difficulty level: {difficulty.upper()} ({guide}). Every question must clearly match this level.

Return ONLY a JSON array (no markdown, no commentary). Each item must be:
{{"question": "...", "options": ["...", "...", "...", "..."], "answer": "B", "explanation": "one or two sentences why"}}
Rules: exactly 4 options per question without letter prefixes; "answer" is a single letter A-D."""

    try:
        if _use_new_genai:
            response = client.models.generate_content(model=MODEL_ID, contents=prompt)
        else:
            response = model.generate_content(prompt)
        quiz_data = _parse_quiz_json(response.text)
        if not quiz_data:
            return jsonify({"error": "Could not read the quiz from the AI. Please try again."}), 500
        return jsonify(quiz_data[:num])
    except Exception as e:
        print("API Quiz Generation Error:", e)
        return jsonify({"error": str(e)}), 500

# ✅ Quiz page (setup -> questions -> result all happen on this page)
@app.route('/quiz_setup')
@app.route('/quiz')
def quiz_setup():
    if "username" not in session:
        return redirect(url_for("login"))
    return render_template("quiz.html", user=session['username'])

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for("index"))

if __name__ == "__main__":
    app.run(debug=True, use_reloader=False)
